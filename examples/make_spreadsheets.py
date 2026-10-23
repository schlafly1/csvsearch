"""Regenerate the sample spreadsheets (needs openpyxl, odfpy, pandas, and xlwt for .xls)."""
import datetime as dt
import os

HERE = os.path.dirname(os.path.abspath(__file__))
ORDERS = [
    ["order_id", "customer", "sku", "qty", "unit_price", "order_date"],
    [1001, "Alice Smith", "A100", 3, 9.99, dt.date(2026, 1, 15)],
    [1002, "Bob Jones", "A200", 1, 19.99, dt.date(2026, 2, 3)],
    [],
    [1003, "Carol Smith", "C400", 12, 7.25, dt.date(2026, 3, 30)],
]
STORES = [["store", "city", "opened"], ["North", "Boston", 1998], ["West", "Denver", 2010]]


def xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    for i, (title, rows) in enumerate([("Orders", ORDERS), ("Stores", STORES)]):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        for r, row in enumerate(rows, 1):
            for c, v in enumerate(row, 1):
                ws.cell(r, c, v)
    wb.save(os.path.join(HERE, "sales.xlsx"))


def xls():
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Inventory")
    datefmt = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    rows = [["sku", "warehouse", "on_hand", "counted"],
            ["A100", "Boston", 250, dt.datetime(2026, 4, 1)],
            ["B300", "Austin", 40, dt.datetime(2026, 4, 2)]]
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            ws.write(r, c, v, datefmt) if isinstance(v, dt.datetime) else ws.write(r, c, v)
    wb.save(os.path.join(HERE, "inventory.xls"))


def ods():
    import pandas as pd
    df = pd.DataFrame(
        [["Boston", "Mass.", 675647, 48.4], ["Denver", "Colo.", 715522, 153.3]],
        columns=["city", "state", "population", "area_sq_mi"],
    )
    df.to_excel(os.path.join(HERE, "cities.ods"), engine="odf", sheet_name="Cities", index=False)


if __name__ == "__main__":
    xlsx(); xls(); ods()
