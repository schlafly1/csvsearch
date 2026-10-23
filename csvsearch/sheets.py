"""Read spreadsheet files (.xlsx, .xls, .ods) as (sheet_name, [(row_number, [str,...])])."""
import datetime as _dt
import os

SPREADSHEET_EXTS = {".xlsx", ".xlsm", ".xls", ".ods"}


def is_spreadsheet(path):
    return os.path.splitext(path)[1].lower() in SPREADSHEET_EXTS


def fmt(v):
    """Format a cell value as text: whole floats without .0, dates as ISO."""
    if v is None:
        return ""
    try:  # pandas NaN / NaT
        if v != v:
            return ""
    except Exception:
        pass
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if hasattr(v, "to_pydatetime"):  # pandas Timestamp
        v = v.to_pydatetime()
    if isinstance(v, _dt.datetime):
        if v.time() == _dt.time(0):
            return v.date().isoformat()
        return v.isoformat(sep=" ")
    if isinstance(v, (_dt.date, _dt.time)):
        return v.isoformat()
    return str(v)


def _trim(row):
    row = list(row)
    while row and row[-1] == "":
        row.pop()
    return row


def _xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = [(i, _trim(fmt(c) for c in r))
                    for i, r in enumerate(ws.iter_rows(values_only=True), start=1)]
            yield ws.title, rows
    finally:
        wb.close()


def _xls(path):
    import xlrd
    wb = xlrd.open_workbook(path)
    for sh in wb.sheets():
        rows = []
        for r in range(sh.nrows):
            vals = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    v = xlrd.xldate.xldate_as_datetime(v, wb.datemode)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    v = bool(v)
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    v = None
                vals.append(fmt(v))
            rows.append((r + 1, _trim(vals)))
        yield sh.name, rows


def _ods(path):
    import pandas as pd
    sheets = pd.read_excel(path, engine="odf", sheet_name=None, header=None, dtype=object)
    for name, df in sheets.items():
        rows = [(i + 1, _trim(fmt(v) for v in rec))
                for i, rec in enumerate(df.itertuples(index=False, name=None))]
        yield name, rows


def read_sheets(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return list(_xlsx(path))
    if ext == ".xls":
        return list(_xls(path))
    if ext == ".ods":
        return list(_ods(path))
    raise ValueError(f"not a spreadsheet: {path}")
